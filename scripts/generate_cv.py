"""Generate Hernán Calvo CV in DOCX and PDF (EN + ES)."""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Cm

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    HRFlowable,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)
from pypdf import PdfReader

OUT_DIR = Path(__file__).resolve().parent.parent / "archivos"
ACCENT = RGBColor(0xEC, 0x18, 0x39)
ACCENT_HEX = colors.HexColor("#ec1839")
DARK = RGBColor(0x30, 0x2E, 0x4D)
MUTED = RGBColor(0x50, 0x4E, 0x70)

CV = {
    "en": {
        "file": "CV Hernan Calvo (English)",
        "name": "HERNÁN CALVO",
        "title": "Semi Senior Full Stack Developer",
        "contact": "Buenos Aires, Argentina  ·  116-212-7899  ·  hernancalvofullstack@gmail.com",
        "links": "Portfolio: portfoliohernancalvo.netlify.app  ·  LinkedIn: linkedin.com/in/hernan-calvo-937a64199  ·  GitHub: github.com/HernanCalvoFullStack  ·  HC Prode: www.hcprode.com.ar",
        "links_pdf": (
            '<a href="https://portfoliohernancalvo.netlify.app/" color="#ec1839"><u>Portfolio</u></a>'
            '  ·  LinkedIn: linkedin.com/in/hernan-calvo-937a64199'
            '  ·  GitHub: github.com/HernanCalvoFullStack'
            '  ·  HC Prode: www.hcprode.com.ar'
        ),
        "summary_title": "PROFESSIONAL SUMMARY",
        "summary": (
            "Semi Senior Full Stack Developer with 2+ years of professional experience building and migrating "
            "production CRM systems, REST APIs and third-party integrations. Strong focus on Laravel, React and "
            "TypeScript. Comfortable owning end-to-end features, database cleanup and system modernization. "
            "Based in Buenos Aires; open to full-time on-site, hybrid or remote roles."
        ),
        "skills_title": "TECHNICAL SKILLS",
        "skills": [
            ("Backend", "Laravel, PHP, Node.js, NestJS, REST APIs, Sequelize, Eloquent, Swagger, Queues & Jobs"),
            ("Frontend", "React, TypeScript, Next.js, JavaScript, HTML5, CSS3, Sass, TailwindCSS"),
            ("Data", "MySQL, PostgreSQL, indexing, relationships, data cleanup, Excel, Power BI"),
            ("Other", "WhatsApp API, Google Wallet, Mercado Pago, Git, Android (personal projects)"),
        ],
        "exp_title": "EXPERIENCE",
        "jobs": [
            {
                "role": "Semi Senior Full Stack Developer",
                "place": "Interassist SA — Travel Assistance",
                "dates": "2024 – Present  ·  Buenos Aires",
                "bullets": [
                    "Migrated CRM/Core from Laravel 4.2 + Blade to Laravel 11 with React REST API; currently upgrading to Laravel 13 + React + TypeScript.",
                    "DBA-style work: cleanup, relationships, indexing and standards; merged duplicate clients with safe SQL ID remapping.",
                    "CRM modules: clients, companies, travel & vehicle providers, case management with automatic status changes.",
                    "Plan quoter, sales issuance, IP access restriction, sector profiles, backend middleware, file imports and promotions.",
                    "Integrations: WhatsApp API, Google Wallet e-vouchers, API Keys and Swagger for external systems.",
                    "Agency and public web quoters with carousel publish/expire via Laravel schedule, jobs and queues.",
                ],
            },
            {
                "role": "Founder & Full Stack Developer",
                "place": "HC Prode — www.hcprode.com.ar",
                "dates": "Personal product  ·  Live",
                "bullets": [
                    "Built a sports prediction platform with Laravel 13 and React + TypeScript.",
                    "Custom and real-world tournaments with company-separated leagues, points engine and Mercado Pago monetization.",
                ],
            },
            {
                "role": "Web & Android App — Expense Splitting",
                "place": "Cuentas Claras",
                "dates": "Personal project  ·  Not published yet",
                "bullets": [
                    "Web and Android app to split expenses for one day or multi-day trips; children link to adults with a payment percentage, per-expense participants, and suggested compensation transfers.",
                ],
            },
        ],
        "edu_title": "EDUCATION",
        "edu": [
            ("Bachelor of Data Science — Universidad Siglo 21 (2023 – Present)", "Distance modality — 18 subjects approved in 1.5 years."),
            ("Mathematics Teaching Degree — Instituto Joaquín V. González (in progress)", "Five finals remaining."),
        ],
        "courses_title": "RELEVANT COURSES",
        "courses": [
            "React and TypeScript — The Complete Guide (+10 projects)",
            "Node.js Bootcamp — MVC and REST APIs",
            "Modern JavaScript — Definitive Guide (+20 projects)",
            "Complete Web Development with HTML5, CSS3, JS, PHP and MySQL",
            "CSS The Complete Guide — Flexbox, CSS Grid, SASS (+20 projects)",
        ],
        "langs_title": "LANGUAGES",
        "langs": "Spanish (Native)  ·  English (Professional working proficiency)",
    },
    "es": {
        "file": "CV Hernan Calvo (Espanol)",
        "name": "HERNÁN CALVO",
        "title": "Desarrollador Full Stack Semi Senior",
        "contact": "Buenos Aires, Argentina  ·  116-212-7899  ·  hernancalvofullstack@gmail.com",
        "links": "Portfolio: portfoliohernancalvo.netlify.app  ·  LinkedIn: linkedin.com/in/hernan-calvo-937a64199  ·  GitHub: github.com/HernanCalvoFullStack  ·  HC Prode: www.hcprode.com.ar",
        "links_pdf": (
            '<a href="https://portfoliohernancalvo.netlify.app/" color="#ec1839"><u>Portfolio</u></a>'
            '  ·  LinkedIn: linkedin.com/in/hernan-calvo-937a64199'
            '  ·  GitHub: github.com/HernanCalvoFullStack'
            '  ·  HC Prode: www.hcprode.com.ar'
        ),
        "summary_title": "PERFIL PROFESIONAL",
        "summary": (
            "Desarrollador Full Stack Semi Senior con más de 2 años de experiencia profesional construyendo y migrando "
            "sistemas CRM en producción, APIs REST e integraciones con terceros. Fuerte foco en Laravel, React y "
            "TypeScript. Cómodo siendo dueño de features de punta a punta, limpieza de bases de datos y modernización "
            "de sistemas. Basado en Buenos Aires; abierto a roles full-time presenciales, híbridos o remotos."
        ),
        "skills_title": "HABILIDADES TÉCNICAS",
        "skills": [
            ("Backend", "Laravel, PHP, Node.js, NestJS, APIs REST, Sequelize, Eloquent, Swagger, Colas & Jobs"),
            ("Frontend", "React, TypeScript, Next.js, JavaScript, HTML5, CSS3, Sass, TailwindCSS"),
            ("Datos", "MySQL, PostgreSQL, indexación, relaciones, limpieza de datos, Excel, Power BI"),
            ("Otros", "WhatsApp API, Google Wallet, Mercado Pago, Git, Android (proyectos personales)"),
        ],
        "exp_title": "EXPERIENCIA",
        "jobs": [
            {
                "role": "Desarrollador Full Stack Semi Senior",
                "place": "Interassist SA — Asistencia al Viajero",
                "dates": "2024 – Actualidad  ·  Buenos Aires",
                "bullets": [
                    "Migré CRM/Core de Laravel 4.2 + Blade a Laravel 11 con API REST en React; actualmente migrando a Laravel 13 + React + TypeScript.",
                    "Administración de BBDD: limpieza, relaciones, indexación y estándares; unificación de clientes duplicados con remapeo seguro de IDs.",
                    "Módulos CRM: clientes, empresas, prestadores de viajeros y vehículos, gestión de casos con cambio automático de estados.",
                    "Cotizador de planes, emisión de ventas, restricción por IP, perfiles, middleware, importaciones y promociones.",
                    "Integraciones: WhatsApp API, Google Wallet, API Keys y Swagger para sistemas externos.",
                    "Cotizadores para agencias y web con carousel programado vía schedule, jobs y colas de Laravel.",
                ],
            },
            {
                "role": "Fundador y Desarrollador Full Stack",
                "place": "HC Prode — www.hcprode.com.ar",
                "dates": "Producto personal  ·  En producción",
                "bullets": [
                    "Plataforma de pronósticos deportivos con Laravel 13 y React + TypeScript.",
                    "Torneos custom y reales por empresa, motor de puntos y monetización con Mercado Pago.",
                ],
            },
            {
                "role": "App web y Android — División de gastos",
                "place": "Cuentas Claras",
                "dates": "Proyecto personal  ·  Aún no publicada",
                "bullets": [
                    "App web y Android para dividir gastos de un día o viajes de varios días; hijos vinculados a adultos con porcentaje, participantes por gasto y compensación sugerida.",
                ],
            },
        ],
        "edu_title": "EDUCACIÓN",
        "edu": [
            ("Licenciatura en Ciencia de Datos — Universidad Siglo 21 (2023 – Actualidad)", "Modalidad a distancia — 18 materias aprobadas en 1.5 años."),
            ("Profesorado de Matemática — Instituto Joaquín V. González (en curso)", "Restan cinco finales."),
        ],
        "courses_title": "CURSOS RELEVANTES",
        "courses": [
            "React y TypeScript — Guía Completa (+10 proyectos)",
            "Bootcamp Node.js — MVC y REST APIs",
            "JavaScript Moderno — Guía Definitiva (+20 proyectos)",
            "Desarrollo Web Completo con HTML5, CSS3, JS, PHP y MySQL",
            "CSS La Guía Completa — Flexbox, CSS Grid, SASS (+20 proyectos)",
        ],
        "langs_title": "IDIOMAS",
        "langs": "Español (Nativo)  ·  Inglés (Nivel profesional de trabajo)",
    },
}


def set_run_font(run, size=10, bold=False, color=DARK, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_hr(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "EC1839")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_heading_line(doc, text, scale=1.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 * scale)
    p.paragraph_format.space_after = Pt(4 * scale)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=True, color=ACCENT)
    add_hr(p)
    return p


def build_docx(data: dict, path: Path, scale: float = 1.08):
    """DOCX spacing mirrored from the PDF scale used for that language."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.2)
        section.bottom_margin = Cm(1.2)
        section.left_margin = Cm(1.45)
        section.right_margin = Cm(1.45)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(2 * scale)
    r = name.add_run(data["name"])
    set_run_font(r, size=20, bold=True, color=DARK)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4 * scale)
    r = title.add_run(data["title"])
    set_run_font(r, size=12, bold=True, color=ACCENT)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(1 * scale)
    r = contact.add_run(data["contact"])
    set_run_font(r, size=9, color=MUTED)

    links = doc.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links.paragraph_format.space_after = Pt(6 * scale)
    r = links.add_run(data["links"])
    set_run_font(r, size=8.5, color=MUTED)

    add_heading_line(doc, data["summary_title"], scale)
    summary = doc.add_paragraph()
    summary.paragraph_format.space_after = Pt(4 * scale)
    summary.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = summary.add_run(data["summary"])
    set_run_font(r, size=10, color=DARK)

    add_heading_line(doc, data["skills_title"], scale)
    for label, value in data["skills"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3 * scale)
        p.paragraph_format.space_before = Pt(0)
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=10, bold=True, color=DARK)
        r2 = p.add_run(value)
        set_run_font(r2, size=10, color=MUTED)

    add_heading_line(doc, data["exp_title"], scale)
    for job in data["jobs"]:
        role = doc.add_paragraph()
        role.paragraph_format.space_before = Pt(7 * scale)
        role.paragraph_format.space_after = Pt(1 * scale)
        r = role.add_run(job["role"])
        set_run_font(r, size=10.5, bold=True, color=DARK)

        place = doc.add_paragraph()
        place.paragraph_format.space_after = Pt(3 * scale)
        r = place.add_run(f"{job['place']}  ·  {job['dates']}")
        set_run_font(r, size=9.5, bold=True, color=ACCENT)

        for bullet in job["bullets"]:
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_after = Pt(2 * scale)
            bp.paragraph_format.space_before = Pt(0)
            bp.paragraph_format.left_indent = Cm(0.45)
            if bp.runs:
                bp.runs[0].text = bullet
                set_run_font(bp.runs[0], size=9.5, color=DARK)
            else:
                r = bp.add_run(bullet)
                set_run_font(r, size=9.5, color=DARK)

    add_heading_line(doc, data["edu_title"], scale)
    for title_txt, detail in data["edu"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4 * scale)
        p.paragraph_format.space_after = Pt(1 * scale)
        r = p.add_run(title_txt)
        set_run_font(r, size=10, bold=True, color=DARK)
        p3 = doc.add_paragraph()
        p3.paragraph_format.space_after = Pt(3 * scale)
        r = p3.add_run(detail)
        set_run_font(r, size=9.5, color=MUTED)

    add_heading_line(doc, data["courses_title"], scale)
    for course in data["courses"]:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_after = Pt(2 * scale)
        bp.paragraph_format.left_indent = Cm(0.45)
        if bp.runs:
            bp.runs[0].text = course
            set_run_font(bp.runs[0], size=9.5, color=DARK)
        else:
            r = bp.add_run(course)
            set_run_font(r, size=9.5, color=DARK)

    add_heading_line(doc, data["langs_title"], scale)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(data["langs"])
    set_run_font(r, size=10, color=DARK)

    doc.save(path)


def _pdf_story(data: dict, scale: float):
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "Name",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        textColor=colors.HexColor("#302e4d"),
        alignment=TA_CENTER,
        spaceAfter=3 * scale,
        leading=21,
    )
    title_style = ParagraphStyle(
        "TitleRole",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11.5,
        textColor=ACCENT_HEX,
        alignment=TA_CENTER,
        spaceAfter=4 * scale,
        leading=14,
    )
    contact_style = ParagraphStyle(
        "Contact",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        textColor=colors.HexColor("#504e70"),
        alignment=TA_CENTER,
        spaceAfter=2 * scale,
        leading=11,
    )
    section_style = ParagraphStyle(
        "Section",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        textColor=ACCENT_HEX,
        spaceBefore=10 * scale,
        spaceAfter=2 * scale,
        leading=13,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        textColor=colors.HexColor("#302e4d"),
        alignment=TA_JUSTIFY,
        leading=13 + 0.9 * (scale - 1),
        spaceAfter=4 * scale,
    )
    skill_style = ParagraphStyle(
        "Skill",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        textColor=colors.HexColor("#302e4d"),
        leading=12.5 + 0.8 * (scale - 1),
        spaceAfter=3 * scale,
    )
    role_style = ParagraphStyle(
        "Role",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        textColor=colors.HexColor("#302e4d"),
        spaceBefore=7 * scale,
        spaceAfter=1 * scale,
        leading=12,
    )
    place_style = ParagraphStyle(
        "Place",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        textColor=ACCENT_HEX,
        spaceAfter=3 * scale,
        leading=11,
    )
    detail_style = ParagraphStyle(
        "Detail",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        textColor=colors.HexColor("#504e70"),
        spaceAfter=4 * scale,
        leading=11,
    )
    bullet_style = ParagraphStyle(
        "Bullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        textColor=colors.HexColor("#302e4d"),
        leading=12 + 0.7 * (scale - 1),
        leftIndent=8,
    )

    story = [
        Paragraph(data["name"], name_style),
        Paragraph(data["title"], title_style),
        Paragraph(data["contact"], contact_style),
        Paragraph(data.get("links_pdf", data["links"]), contact_style),
        Spacer(1, 6 * scale),
    ]

    def section(title):
        story.append(Paragraph(title, section_style))
        story.append(
            HRFlowable(
                width="100%",
                thickness=1,
                color=ACCENT_HEX,
                spaceBefore=0,
                spaceAfter=5 * scale,
            )
        )

    section(data["summary_title"])
    story.append(Paragraph(data["summary"], body_style))

    section(data["skills_title"])
    for label, value in data["skills"]:
        story.append(Paragraph(f"<b>{label}:</b> {value}", skill_style))

    section(data["exp_title"])
    for job in data["jobs"]:
        story.append(Paragraph(job["role"], role_style))
        story.append(Paragraph(f"{job['place']}  ·  {job['dates']}", place_style))
        items = [
            ListItem(Paragraph(bullet, bullet_style), leftIndent=8, bulletColor=ACCENT_HEX)
            for bullet in job["bullets"]
        ]
        story.append(
            ListFlowable(
                items,
                bulletType="bullet",
                start="•",
                leftIndent=12,
                spaceBefore=0,
                spaceAfter=4 * scale,
                bulletFontSize=9,
            )
        )

    section(data["edu_title"])
    for title_txt, detail in data["edu"]:
        story.append(Paragraph(f"<b>{title_txt}</b>", skill_style))
        story.append(Paragraph(detail, detail_style))

    section(data["courses_title"])
    course_items = [
        ListItem(Paragraph(c, bullet_style), leftIndent=8, bulletColor=ACCENT_HEX)
        for c in data["courses"]
    ]
    story.append(
        ListFlowable(
            course_items,
            bulletType="bullet",
            start="•",
            leftIndent=12,
            spaceBefore=0,
            spaceAfter=4 * scale,
            bulletFontSize=9,
        )
    )

    section(data["langs_title"])
    story.append(Paragraph(data["langs"], body_style))
    return story


def _pdf_page_count(data: dict, scale: float) -> int:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=1.45 * cm,
        rightMargin=1.45 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
    )
    doc.build(_pdf_story(data, scale))
    return len(PdfReader(buf).pages)


def fit_pdf_scale(data: dict, lo: float = 1.0, hi: float = 1.55) -> float:
    """Largest spacing scale that still fits on one A4 page."""
    if _pdf_page_count(data, lo) > 1:
        return lo
    best = lo
    for _ in range(14):
        mid = (lo + hi) / 2
        if _pdf_page_count(data, mid) <= 1:
            best = mid
            lo = mid
        else:
            hi = mid
    # slight safety margin so rendering quirks don't spill
    return max(1.0, best * 0.97)


def build_pdf(data: dict, path: Path, scale: float | None = None) -> float:
    if scale is None:
        scale = fit_pdf_scale(data)
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=1.45 * cm,
        rightMargin=1.45 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
    )
    doc.build(_pdf_story(data, scale))
    return scale


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for lang, data in CV.items():
        base = OUT_DIR / data["file"]
        docx_path = base.with_suffix(".docx")
        pdf_path = base.with_suffix(".pdf")
        scale = fit_pdf_scale(data)
        build_docx(data, docx_path, scale=scale)
        build_pdf(data, pdf_path, scale=scale)
        pages = len(PdfReader(str(pdf_path)).pages)
        print(f"Created: {docx_path.name}")
        print(f"Created: {pdf_path.name}  (scale={scale:.3f}, pages={pages})")


if __name__ == "__main__":
    main()
